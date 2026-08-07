#!/usr/bin/env python3
"""
glm head-to-head
Compares sparse imagery datasets across two S3 buckets using SSO authentication,
a configuration file, dynamic date-based prefix routing, and deep NetCDF metadata comparison
powered by xarray (META-PAVE engine), outputting to a Landscape Word Doc with 1x3 Spatial Plots.
"""
import base_utils
import aws_utils
import meta_utils
import json
import sys
import argparse
import os
import csv
import difflib
import tempfile
from datetime import datetime, timezone

try:
    import s3fs
    HAS_S3FS = True
except ImportError:
    HAS_S3FS = False

try:
    import docx
    from docx.shared import Inches, Pt
    from docx.enum.section import WD_ORIENT
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

try:
    import numpy as np
    import matplotlib
    matplotlib.use('Agg') # Headless plotting
    import matplotlib.pyplot as plt
    HAS_MPL = True
except ImportError:
    HAS_MPL = False

try:
    import cartopy.crs as ccrs
    import cartopy.feature as cfeature
    HAS_CARTOPY = True
except ImportError:
    HAS_CARTOPY = False


# =============================================================================
# PLOTTING & DOCX UTILITIES
# =============================================================================
def generate_diff_plot(var_name, v1_data, v2_data, name1, name2):
    """Generates a graphical plot of the numeric differences between two arrays."""
    if not HAS_MPL: return None
    try:
        v1_flat = np.asanyarray(v1_data).flatten()
        v2_flat = np.asanyarray(v2_data).flatten()

        if v1_flat.size != v2_flat.size:
            return None

        diff = v1_flat - v2_flat
        valid_mask = ~(np.isnan(diff) | np.isinf(diff))
        diff = diff[valid_mask]

        if len(diff) == 0 or np.all(diff == 0):
            return None

        plt.figure(figsize=(6.5, 3.5))

        if len(diff) > 5000:
            plt.hist(diff, bins=50, color='tab:red', alpha=0.7)
            plt.title(f"{var_name} Difference Distribution")
            plt.xlabel(f"Difference ({name1} - {name2})")
            plt.ylabel("Frequency")
        else:
            plt.plot(diff, marker='o', markersize=3, linestyle='none', color='tab:red', alpha=0.6)
            plt.title(f"{var_name} Mismatched Values")
            plt.xlabel("Index")
            plt.ylabel(f"Difference ({name1} - {name2})")
            plt.axhline(0, color='black', linewidth=0.8, linestyle='--')

        plt.tight_layout()
        tmp_fd, tmp_path = tempfile.mkstemp(suffix=".png")
        os.close(tmp_fd)
        plt.savefig(tmp_path, dpi=150)
        plt.close()
        return tmp_path
    except Exception as e:
        print(f"Warning: Failed to generate diff plot: {e}")
        return None

def generate_spatial_plot(entity_name, lat1, lon1, lat2, lon2, name1, name2):
    """Generates a 1x3 spatial subplot mapping Source 1, Differences, and Source 2."""
    if not HAS_MPL: return None
    try:
        mask1 = ~(np.isnan(lat1) | np.isnan(lon1))
        lat1, lon1 = lat1[mask1], lon1[mask1]

        mask2 = ~(np.isnan(lat2) | np.isnan(lon2))
        lat2, lon2 = lat2[mask2], lon2[mask2]

        if len(lat1) == 0 and len(lat2) == 0:
            return None

        fig = plt.figure(figsize=(12, 4))

        all_lats = np.concatenate([lat1, lat2])
        all_lons = np.concatenate([lon1, lon2])
        if len(all_lats) > 0:
            min_lat, max_lat = np.min(all_lats) - 2, np.max(all_lats) + 2
            min_lon, max_lon = np.min(all_lons) - 2, np.max(all_lons) + 2
        else:
            min_lat, max_lat, min_lon, max_lon = -90, 90, -180, 180

        extent = [min_lon, max_lon, min_lat, max_lat]
        axes = []

        if HAS_CARTOPY:
            for i in range(1, 4):
                ax = fig.add_subplot(1, 3, i, projection=ccrs.PlateCarree())
                ax.add_feature(cfeature.COASTLINE, linewidth=0.8)
                ax.add_feature(cfeature.BORDERS, linewidth=0.5, linestyle=':')
                ax.add_feature(cfeature.STATES, linewidth=0.3, linestyle=':')
                gl = ax.gridlines(draw_labels=True, linestyle='--', alpha=0.5)
                gl.top_labels = False
                gl.right_labels = False
                if i > 1: gl.left_labels = False
                ax.set_extent(extent, crs=ccrs.PlateCarree())
                axes.append(ax)
            scatter_kwargs = {'transform': ccrs.PlateCarree()}
        else:
            for i in range(1, 4):
                ax = fig.add_subplot(1, 3, i)
                ax.grid(True, linestyle='--', alpha=0.5)
                ax.set_xlim([min_lon, max_lon])
                ax.set_ylim([min_lat, max_lat])
                if i == 1: ax.set_ylabel("Latitude")
                ax.set_xlabel("Longitude")
                axes.append(ax)
            scatter_kwargs = {}

        # Left
        axes[0].set_title(f"{name1}")
        if len(lat1) > 0:
            axes[0].scatter(lon1, lat1, color='tab:blue', s=4, alpha=0.6, marker='o', **scatter_kwargs)

        # Center
        axes[1].set_title("Overlay / Differences")
        if len(lat1) > 0:
            axes[1].scatter(lon1, lat1, color='tab:blue', s=4, alpha=0.6, marker='o', label=name1, **scatter_kwargs)
        if len(lat2) > 0:
            axes[1].scatter(lon2, lat2, color='tab:red', s=4, alpha=0.6, marker='x', label=name2, **scatter_kwargs)
        axes[1].legend(loc="upper right", fontsize='small')

        # Right
        axes[2].set_title(f"{name2}")
        if len(lat2) > 0:
            axes[2].scatter(lon2, lat2, color='tab:red', s=4, alpha=0.6, marker='x', **scatter_kwargs)

        fig.suptitle(f"{entity_name.capitalize()}s Spatial Distribution", fontsize=14, y=1.05)
        plt.tight_layout()

        tmp_fd, tmp_path = tempfile.mkstemp(suffix=".png")
        os.close(tmp_fd)
        plt.savefig(tmp_path, dpi=150, bbox_inches='tight')
        plt.close(fig)
        return tmp_path
    except Exception as e:
        print(f"Warning: Failed to generate spatial plot: {e}")
        return None

def add_diff_runs_to_cells(cell1, cell2, text1, text2, font_size=None):
    """Highlights differences between text1 and text2 by bolding them in docx cells."""
    text1, text2 = str(text1), str(text2)

    p1 = cell1.paragraphs[0]
    p1.text = ""
    p2 = cell2.paragraphs[0]
    p2.text = ""

    matcher = difflib.SequenceMatcher(None, text1, text2)
    for tag, i1, i2, j1, j2 in matcher.get_opcodes():
        r1, r2 = None, None
        if tag == 'equal':
            r1 = p1.add_run(text1[i1:i2])
            r2 = p2.add_run(text2[j1:j2])
        elif tag == 'replace':
            r1 = p1.add_run(text1[i1:i2])
            r1.bold = True
            r2 = p2.add_run(text2[j1:j2])
            r2.bold = True
        elif tag == 'delete':
            r1 = p1.add_run(text1[i1:i2])
            r1.bold = True
        elif tag == 'insert':
            r2 = p2.add_run(text2[j1:j2])
            r2.bold = True

        if font_size:
            if r1: r1.font.size = font_size
            if r2: r2.font.size = font_size


# =============================================================================
# DATASET COMPARATOR CORE
# =============================================================================
def resolve_placeholders(template_str, target_date):
    """Replaces date tags in a string with values from the target_date."""
    if not template_str:
        return ""

    replacements = {
        "<YEAR>": target_date.strftime("%Y"),
        "<YYYY>": target_date.strftime("%Y"),
        "<mon>": target_date.strftime("%b").lower(),
        "<MM>": target_date.strftime("%m"),
        "<DD>": target_date.strftime("%d"),
        "<YYYYMMDD>": target_date.strftime("%Y%m%d"),
        "<DOY>": target_date.strftime("%j"),
        "<HH>": target_date.strftime("%H")
    }

    resolved_str = template_str
    for tag, value in replacements.items():
        resolved_str = resolved_str.replace(tag, value)

    return resolved_str

def compare_datasets(s3_client, ds1_cfg, ds2_cfg, target_date, raw_date_str, ext_filter=None, logger=None):
    """Compares two S3 buckets/prefixes and prints the differences."""
    name1 = resolve_placeholders(ds1_cfg.get('name', 'Source 1'), target_date)
    b1 = ds1_cfg.get('bucket')
    p1 = resolve_placeholders(ds1_cfg.get('prefix', ''), target_date)

    name2 = resolve_placeholders(ds2_cfg.get('name', 'Source 2'), target_date)
    b2 = ds2_cfg.get('bucket')
    p2 = resolve_placeholders(ds2_cfg.get('prefix', ''), target_date)

    objs1 = aws_utils.get_s3_objects(s3_client, b1, p1, ext_filter, raw_date_str, logger)
    objs2 = aws_utils.get_s3_objects(s3_client, b2, p2, ext_filter, raw_date_str, logger)

    if objs1 is None or objs2 is None:
        if logger: logger.error("Comparison aborted due to bucket access errors.")
        sys.exit(1)

    set1, set2 = set(objs1.keys()), set(objs2.keys())

    only_in_1 = set1 - set2
    only_in_2 = set2 - set1
    common = set1.intersection(set2)

    identical = []
    mismatched = []
    meta_diffs = {}
    all_raw_issues = {} # Store for Word Doc Generation

    auditor = meta_utils.MetadataAuditor()

    csv_report_file = f"reports/glm_comparison_report_{raw_date_str}.csv"
    doc_report_file = f"reports/glm_comparison_report_{raw_date_str}.docx"

    try:
        with open(csv_report_file, 'w', newline='') as f:
            writer = csv.writer(f)
            writer.writerow(["Normalized Key", f"{name1} Key", f"{name2} Key", "Attribute", "Status", f"{name1} Value", f"{name2} Value"])
    except Exception as e:
        if logger: logger.error(f"Could not create CSV report: {e}")

    credentials = s3_client._request_signer._credentials
    fs = None
    if HAS_S3FS and credentials:
        fs = s3fs.S3FileSystem(
            key=credentials.access_key,
            secret=credentials.secret_key,
            token=credentials.token,
            client_kwargs={'region_name': s3_client.meta.region_name}
        )

    for key in common:
        file1, file2 = objs1[key], objs2[key]

        if file1['Size'] == file2['Size'] and file1['ETag'] == file2['ETag']:
            identical.append(key)
        else:
            mismatched.append({
                'norm_key': key,
                'b1_key': file1['full_key'], 'b2_key': file2['full_key'],
                'b1_size': file1['Size'], 'b2_size': file2['Size'],
                'b1_etag': file1['ETag'], 'b2_etag': file2['ETag']
            })

            if file1['full_key'].lower().endswith('.nc'):
                if logger: logger.info(f"Streaming mismatched file for deep metadata & data validation: {key}")

                uri1 = f"s3://{b1}/{file1['full_key']}"
                uri2 = f"s3://{b2}/{file2['full_key']}"

                try:
                    # Pass the plotting functions into the auditor as callbacks!
                    raw_issues = auditor.audit_file_pair(
                        uri1, uri2, fs, name1, name2,
                        diff_plot_cb=generate_diff_plot,
                        spatial_plot_cb=generate_spatial_plot
                    )
                    all_raw_issues[key] = raw_issues

                    filtered_diffs = []
                    with open(csv_report_file, 'a', newline='') as f:
                        writer = csv.writer(f)
                        for issue in raw_issues:
                            if issue['Status'] == 'PLOT':
                                filtered_diffs.append(f"[{issue['Status']:<7}] {issue['Attribute']} | Spatial plot generated for {name1} vs {name2}.")
                                continue

                            if issue['Status'] != 'IGNORE':
                                filtered_diffs.append(
                                    f"[{issue['Status']:<7}] {issue['Attribute']} | {name1}: {issue['Source1']} | {name2}: {issue['Source2']}"
                                )
                                writer.writerow([
                                    key, file1['full_key'], file2['full_key'],
                                    issue['Attribute'], issue['Status'],
                                    issue['Source1'], issue['Source2']
                                ])

                    if filtered_diffs:
                        meta_diffs[key] = filtered_diffs
                    else:
                        meta_diffs[key] = ["No internal differences found. Size/ETag mismatch may be due to NetCDF chunking differences."]
                except Exception as e:
                    meta_diffs[key] = [f"[ERROR  ] S3 Stream/Compare failed: {e}"]

    # --- Print Summary ---
    print("\n" + "="*80)
    print(f" COMPARISON RESULTS (Date Filter: {raw_date_str})")
    print("="*80)
    print(f"{name1}: s3://{b1}/{p1}")
    print(f"{name2}: s3://{b2}/{p2}")
    if ext_filter: print(f"Filtered by: {ext_filter}")
    print("-" * 80)

    print(f"Total files in {name1}: {len(objs1)}")
    print(f"Total files in {name2}: {len(objs2)}")
    print(f"Files ONLY in {name1}:  {len(only_in_1)}")
    print(f"Files ONLY in {name2}:  {len(only_in_2)}")
    print(f"Files in BOTH (Names):   {len(common)}")
    print(f"  -> Identical Content:  {len(identical)}")
    print(f"  -> Mismatched Content: {len(mismatched)}")
    print("="*80)

    # --- Generate Streamlined Word Document Report (Plots Only) ---
    if HAS_DOCX and mismatched:
        if logger: logger.info("Generating Word Document report with spatial plots...")
        try:
            doc = docx.Document()

            # ---> SET TO LANDSCAPE & ADJUST MARGINS <---
            for section in doc.sections:
                new_width, new_height = section.page_height, section.page_width
                section.orientation = WD_ORIENT.LANDSCAPE
                section.page_width = new_width
                section.page_height = new_height

                section.top_margin = Inches(1)
                section.bottom_margin = Inches(1)
                section.left_margin = Inches(1)
                section.right_margin = Inches(1)

            doc.add_heading('GLM Dataset Spatial Comparison Report', 0)

            doc.add_heading('Summary', level=1)
            doc.add_paragraph(f"Date Filter: {raw_date_str}")
            doc.add_paragraph(f"{name1}: s3://{b1}/{p1}")
            doc.add_paragraph(f"{name2}: s3://{b2}/{p2}")
            doc.add_paragraph(f"Mismatched Files Analysed: {len(mismatched)}")

            doc.add_heading('Mismatched Files Spatial Analysis', level=1)
            plots_to_delete = []

            for m in mismatched:
                k = m['norm_key']
                issues = all_raw_issues.get(k, [])

                # Filter strictly for spatial plots to keep report streamlined
                spatial_issues = [i for i in issues if i.get('Status') == 'PLOT' and 'GLM_Spatial' in i.get('Attribute', '')]

                if spatial_issues:
                    doc.add_heading(f"{m['b1_key']}", level=2)
                    for issue in spatial_issues:
                        plot_path = issue.get('Plot')
                        if plot_path and os.path.exists(plot_path):
                            p = doc.add_paragraph()
                            p.add_run(f"{issue.get('Attribute').split(':')[-1]}").bold = True

                            # Expand to 9 inches wide to fully utilize the landscape layout
                            doc.add_picture(plot_path, width=Inches(9.0))
                            plots_to_delete.append(plot_path)

            doc.save(doc_report_file)
            print(f"\nDetailed reports generated:\n  - {csv_report_file} (Metadata Diff)\n  - {doc_report_file} (Spatial Plots)")

            for p in plots_to_delete:
                try: os.unlink(p)
                except: pass

        except Exception as e:
            if logger: logger.error(f"Failed to generate Word doc: {e}")
            print(f"\nDetailed CSV report generated: {csv_report_file}")
    else:
        if mismatched:
            print(f"\nDetailed CSV report generated: {csv_report_file}")
            if not HAS_DOCX or not HAS_MPL:
                print("(Note: Install 'python-docx' and 'matplotlib' to generate a graphic Word document report)")


if __name__ == "__main__":
    logger = base_utils.Logger(level="INFO")

    parser = argparse.ArgumentParser(description="Compare imagery datasets with dynamic date placeholders.")
    parser.add_argument(
        "--date",
        type=str,
        default=datetime.now(timezone.utc).strftime("%Y%j%H"),
        help="Target date/time in YYYYDOY[HH[MM[SS]]] format (defaults to current UTC hour)."
    )
    args = parser.parse_args()

    raw_date = args.date
    try:
        if len(raw_date) == 7: target_date = datetime.strptime(raw_date, "%Y%j")
        elif len(raw_date) == 9: target_date = datetime.strptime(raw_date, "%Y%j%H")
        elif len(raw_date) == 11: target_date = datetime.strptime(raw_date, "%Y%j%H%M")
        elif len(raw_date) == 13: target_date = datetime.strptime(raw_date, "%Y%j%H%M%S")
        else: raise ValueError()
    except ValueError:
        logger.error(f"Invalid date format for '{raw_date}'. Please use YYYYDOY, YYYYDOYHH, YYYYDOYHHMM, or YYYYDOYHHMMSS.")
        sys.exit(1)

    config_file = "config.json"
    try:
        with open(config_file, 'r') as f:
            config = json.load(f)
    except Exception as e:
        logger.error(f"Failed to read {config_file}: {e}")
        sys.exit(1)

    data_sources = config.get("data_sources", {})

    dataset_1 = data_sources.get("dataset_1")
    dataset_2 = data_sources.get("dataset_2")
    ext_filter = data_sources.get("extension_filter")

    if not dataset_1 or not dataset_2 or 'bucket' not in dataset_1 or 'bucket' not in dataset_2:
        logger.error("Configuration must contain 'dataset_1' and 'dataset_2' inside 'data_sources', both with a 'bucket' key.")
        sys.exit(1)

    logger.info("Initializing SSO Connection...")
    s3_connector = aws_utils.S3SSOConnector(config_path=config_file)

    logger.info(f"Starting dataset comparison for date filter: {raw_date}...")
    compare_datasets(
        s3_client=s3_connector.s3_client,
        ds1_cfg=dataset_1,
        ds2_cfg=dataset_2,
        target_date=target_date,
        raw_date_str=raw_date,
        ext_filter=ext_filter,
        logger=logger
    )
